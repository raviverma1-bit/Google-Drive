#!/usr/bin/env python3
"""
Render-deployable web app: upload an Excel file listing (identifier, Drive
folder link) pairs, and it bulk-downloads each folder (must be public /
"anyone with the link"), naming each output folder after your identifier,
then zips everything for download.

Deploy on Render as a "Web Service":
  - Build command:  pip install -r requirements.txt
  - Start command:  gunicorn app:app --timeout 600

Excel format expected (first row = headers, any column order):
  Identifier | Link
  KA03AJ5970 | https://drive.google.com/drive/folders/1BYctOnok9PaXFIAz-eh42EPNIOH5882d
  ...

Header matching is case-insensitive and tolerant of these names:
  Identifier: "identifier", "id", "name", "reg no", "registration"
  Link:       "link", "url", "folder link", "drive link"

NOTE: Render's free/starter instances have ephemeral disk and can spin
down after inactivity. Download the zip promptly after the job finishes;
don't rely on this as permanent storage. For very large or long-running
jobs, prefer a paid instance or a Render Background Worker/Cron Job.
"""

import os
import re
import threading
import shutil
import zipfile
from datetime import datetime

from flask import Flask, jsonify, send_file, request, render_template_string
from openpyxl import load_workbook, Workbook

try:
    import gdown
except ImportError:
    gdown = None

app = Flask(__name__)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_PATH = os.path.join(BASE_DIR, "uploaded_links.xlsx")
DOWNLOAD_DIR = os.path.join(BASE_DIR, "drive_downloads")
ZIP_PATH = os.path.join(BASE_DIR, "drive_downloads.zip")

IDENTIFIER_HEADERS = {"identifier", "id", "name", "reg no", "registration", "reg. no", "code"}
LINK_HEADERS = {"link", "url", "folder link", "drive link", "folder url", "drive url"}

state = {
    "status": "idle",            # idle | parsed | running | zipping | done | error
    "total": 0,
    "completed": 0,
    "current_identifier": None,
    "log": [],
    "zip_ready": False,
    "started_at": None,
    "finished_at": None,
    "rows_preview": [],          # first few parsed rows, for confirmation
    "parse_error": None,
}
state_lock = threading.Lock()

# Parsed (identifier, link) pairs live here between /upload and /start
parsed_rows = []
rows_lock = threading.Lock()


def log(msg: str):
    line = f"[{datetime.utcnow().strftime('%H:%M:%S')}] {msg}"
    with state_lock:
        state["log"].append(line)
        state["log"] = state["log"][-500:]
    print(line, flush=True)


def sanitize_name(name: str) -> str:
    name = str(name).strip()
    name = re.sub(r"[^\w\-. ]", "_", name)
    return name[:100] if name else "unnamed"


def extract_folder_id(url: str) -> str:
    match = re.search(r"/folders/([a-zA-Z0-9_-]+)", str(url))
    if match:
        return match.group(1)
    # Also accept a bare folder ID pasted directly
    if re.fullmatch(r"[a-zA-Z0-9_-]{10,}", str(url).strip()):
        return str(url).strip()
    raise ValueError(f"Could not find a folder ID in: {url}")


def parse_excel(path: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The spreadsheet appears to be empty.")

    header_row = rows[0]
    headers = [str(h).strip().lower() if h is not None else "" for h in header_row]

    id_col = next((i for i, h in enumerate(headers) if h in IDENTIFIER_HEADERS), None)
    link_col = next((i for i, h in enumerate(headers) if h in LINK_HEADERS), None)

    if id_col is None or link_col is None:
        raise ValueError(
            "Couldn't find both an identifier column and a link column. "
            f"Headers found: {header_row}. "
            "Expected something like 'Identifier' and 'Link'."
        )

    results = []
    for row in rows[1:]:
        if row is None or all(v is None for v in row):
            continue
        identifier_raw = row[id_col]
        link_raw = row[link_col]
        if not identifier_raw or not link_raw:
            continue
        results.append({
            "identifier": sanitize_name(identifier_raw),
            "link": str(link_raw).strip(),
        })

    if not results:
        raise ValueError("No valid (identifier, link) rows found after the header row.")

    return results


def run_job():
    if gdown is None:
        with state_lock:
            state["status"] = "error"
        log("gdown is not installed. Check requirements.txt / build logs.")
        return

    with rows_lock:
        rows = list(parsed_rows)

    if not rows:
        log("No rows to process. Upload a spreadsheet first.")
        with state_lock:
            state["status"] = "error"
        return

    with state_lock:
        state["status"] = "running"
        state["total"] = len(rows)
        state["completed"] = 0
        state["log"] = []
        state["zip_ready"] = False
        state["started_at"] = datetime.utcnow().isoformat()
        state["finished_at"] = None

    if os.path.exists(DOWNLOAD_DIR):
        shutil.rmtree(DOWNLOAD_DIR)
    os.makedirs(DOWNLOAD_DIR, exist_ok=True)

    for row in rows:
        identifier = row["identifier"]
        link = row["link"]
        dest = os.path.join(DOWNLOAD_DIR, identifier)
        os.makedirs(dest, exist_ok=True)

        with state_lock:
            state["current_identifier"] = identifier

        log(f"Downloading '{identifier}' <- {link}")
        try:
            folder_id = extract_folder_id(link)
            gdown.download_folder(
                url=f"https://drive.google.com/drive/folders/{folder_id}",
                output=dest,
                quiet=True,
                use_cookies=False,
            )
            log(f"  OK: {identifier}")
        except Exception as e:
            log(f"  FAILED: {identifier} -> {e}")

        with state_lock:
            state["completed"] += 1

    log("All rows processed. Creating zip archive ...")
    with state_lock:
        state["status"] = "zipping"

    if os.path.exists(ZIP_PATH):
        os.remove(ZIP_PATH)

    with zipfile.ZipFile(ZIP_PATH, "w", zipfile.ZIP_DEFLATED) as zf:
        for root, _, files in os.walk(DOWNLOAD_DIR):
            for f in files:
                full_path = os.path.join(root, f)
                arcname = os.path.relpath(full_path, DOWNLOAD_DIR)
                zf.write(full_path, arcname)

    log("Zip archive ready.")
    with state_lock:
        state["status"] = "done"
        state["zip_ready"] = True
        state["finished_at"] = datetime.utcnow().isoformat()


PAGE = """
<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>Drive Folder Bulk Downloader</title>
  <style>
    body { font-family: system-ui, sans-serif; max-width: 760px; margin: 40px auto; padding: 0 16px; }
    button { padding: 10px 18px; font-size: 15px; cursor: pointer; margin-right: 8px; }
    button:disabled { opacity: 0.5; cursor: not-allowed; }
    input[type=file] { margin: 12px 0; }
    #log { background: #111; color: #0f0; font-family: monospace; font-size: 12px;
           padding: 12px; height: 320px; overflow-y: auto; white-space: pre-wrap; border-radius: 6px; }
    .bar-bg { background: #eee; border-radius: 6px; height: 18px; overflow: hidden; margin: 12px 0; }
    .bar { background: #4caf50; height: 100%; width: 0%; transition: width 0.3s; }
    a.download { display: inline-block; margin-top: 12px; }
    table { border-collapse: collapse; margin-top: 10px; font-size: 13px; }
    td, th { border: 1px solid #ccc; padding: 4px 8px; text-align: left; }
    .hint { color: #666; font-size: 13px; }
    .error { color: #b00020; }
  </style>
</head>
<body>
  <h2>Drive Folder Bulk Downloader</h2>
  <p class="hint">
    Upload an Excel file (.xlsx) with two columns: an <b>Identifier</b> column
    (used to name each output folder) and a <b>Link</b> column (the Google Drive
    folder link — must be shared as "anyone with the link").
    <a href="/template">Download a template</a>.
  </p>

  <form id="uploadForm">
    <input type="file" id="fileInput" name="file" accept=".xlsx" required>
    <button type="submit">Upload &amp; Parse</button>
  </form>
  <div id="parseResult"></div>

  <hr>
  <button id="startBtn" onclick="start()" disabled>Start Download</button>
  <div class="bar-bg"><div class="bar" id="bar"></div></div>
  <div id="statusText">Idle. Upload a spreadsheet to begin.</div>
  <a class="download" id="dlLink" href="/download" style="display:none;">⬇ Download ZIP</a>
  <h3>Log</h3>
  <div id="log"></div>

<script>
const uploadForm = document.getElementById('uploadForm');
uploadForm.addEventListener('submit', function(e) {
  e.preventDefault();
  const fd = new FormData();
  fd.append('file', document.getElementById('fileInput').files[0]);
  fetch('/upload', { method: 'POST', body: fd })
    .then(r => r.json())
    .then(res => {
      const el = document.getElementById('parseResult');
      if (res.ok) {
        el.innerHTML = `<p>Parsed <b>${res.count}</b> rows. Preview:</p>` + renderPreview(res.preview);
        document.getElementById('startBtn').disabled = false;
      } else {
        el.innerHTML = `<p class="error">Error: ${res.error}</p>`;
        document.getElementById('startBtn').disabled = true;
      }
    });
});

function renderPreview(rows) {
  let html = '<table><tr><th>Identifier</th><th>Link</th></tr>';
  rows.forEach(r => { html += `<tr><td>${r.identifier}</td><td>${r.link}</td></tr>`; });
  html += '</table>';
  return html;
}

function start() {
  document.getElementById('startBtn').disabled = true;
  fetch('/start', { method: 'POST' }).then(poll);
}
function poll() {
  fetch('/status').then(r => r.json()).then(s => {
    const pct = s.total ? Math.round((s.completed / s.total) * 100) : 0;
    document.getElementById('bar').style.width = pct + '%';
    document.getElementById('statusText').textContent =
      `Status: ${s.status} — ${s.completed}/${s.total} folders` +
      (s.current_identifier ? ` (current: ${s.current_identifier})` : '');
    document.getElementById('log').textContent = s.log.join('\\n');
    document.getElementById('log').scrollTop = document.getElementById('log').scrollHeight;
    if (s.zip_ready) {
      document.getElementById('dlLink').style.display = 'inline-block';
    }
    if (s.status === 'running' || s.status === 'zipping') {
      setTimeout(poll, 1500);
    } else if (s.status !== 'idle') {
      document.getElementById('startBtn').disabled = false;
    }
  });
}
window.onload = poll;
</script>
</body>
</html>
"""


@app.route("/")
def index():
    return render_template_string(PAGE)


@app.route("/template")
def template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Links"
    ws.append(["Identifier", "Link"])
    ws.append(["KA03AJ5970", "https://drive.google.com/drive/folders/1BYctOnok9PaXFIAz-eh42EPNIOH5882d"])
    ws.append(["MH12AB1234", "https://drive.google.com/drive/folders/1SfjQsgbIoNHTdggX2_kXFtOzkQ91gvZX"])
    tmp_path = os.path.join(BASE_DIR, "template.xlsx")
    wb.save(tmp_path)
    return send_file(tmp_path, as_attachment=True, download_name="drive_links_template.xlsx")


@app.route("/upload", methods=["POST"])
def upload():
    file = request.files.get("file")
    if not file:
        return jsonify({"ok": False, "error": "No file received."}), 400

    file.save(UPLOAD_PATH)

    try:
        rows = parse_excel(UPLOAD_PATH)
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400

    with rows_lock:
        parsed_rows.clear()
        parsed_rows.extend(rows)

    with state_lock:
        state["status"] = "parsed"
        state["total"] = len(rows)
        state["completed"] = 0
        state["rows_preview"] = rows[:5]
        state["parse_error"] = None

    return jsonify({"ok": True, "count": len(rows), "preview": rows[:5]})


@app.route("/start", methods=["POST"])
def start():
    with state_lock:
        if state["status"] in ("running", "zipping"):
            return jsonify({"ok": False, "message": "Job already running"}), 409
    with rows_lock:
        if not parsed_rows:
            return jsonify({"ok": False, "message": "Upload a spreadsheet first"}), 400
    thread = threading.Thread(target=run_job, daemon=True)
    thread.start()
    return jsonify({"ok": True})


@app.route("/status")
def status():
    with state_lock:
        return jsonify(dict(state))


@app.route("/download")
def download():
    if not os.path.exists(ZIP_PATH):
        return "Zip not ready yet.", 404
    return send_file(ZIP_PATH, as_attachment=True, download_name="drive_downloads.zip")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
